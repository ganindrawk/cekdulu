"""Bangun spreadsheet katalog awal CekDulu (kategori kecantikan) dari hasil riset.

Kolom link_affiliate dikosongkan untuk diisi Kan (harus dari akun affiliate-nya).
Jalankan: python3 build_catalog_xlsx.py
Output: ~/workspace/your_files/katalog-cekdulu-beauty.xlsx
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADERS = [
    "nama_produk",
    "marketplace",
    "url_asli",
    "harga",
    "rating",
    "jumlah_ulasan",
    "link_affiliate",
    "url_gambar",
    "deskripsi",
]

# (nama, url_asli, harga_min, rating, jumlah_ulasan)
PRODUCTS = [
    ("EMINA Facial Wash Niacinamide Oxy Ceramide Brightening",
     "https://shopee.co.id/-NEW!-EMINA-Facial-Wash-Niacinamide-Oxy-Ceramide-Brightening-Sabun-Cuci-Muka-Mencerahkan-Facial-Wash-Glowing-Pembersih-Wajah-Lembab-Tanpa-Ketarik-Oil-Control-Cleanser-Untuk-Segala-Jenis-Kulit-Bright-Stuff-i.63983008.2202837872",
     21500, 4.9, 223300),
    ("ELFORMULA Intensive Peeling Solution 35gr",
     "https://shopee.co.id/ELFORMULA-Intensive-Peeling-Solution-35-gr-Peeling-Serum-Dark-Spot-Glass-Skin-Eksfoliasi-Wajah-Sel-Kulit-Mati-AHA-BHA-PHA-Glowing-Smooth-Skin-Mencerahkan-Kulit-Halus-Blackhead-Brightening-Cocok-Kulit-Berjerawat-Komedo-acne-Retinol-i.600719961.15304662949",
     109000, 4.9, 547000),
    ("SKINTIFIC Panthenol Gentle Gel Cleanser 120ml",
     "https://shopee.co.id/SKINTIFIC-Panthenol-Gentle-Gel-Cleanser-120ml-Sabun-Pembersih-Wajah-Gentle-Ultimate-Moisturizing-Pembersihan-Pori-pori-Kulit-Berminyak-Sabun-Cuci-Muka-Oil-Control-Deep-Clean-Facial-Wash-Acne-Prone-Skin-3pcs-Bundle-i.380266264.20790142707",
     63000, 4.9, 202500),
    ("SCARLETT Peeling Gel Peeling So Good",
     "https://shopee.co.id/SCARLETT-Peeling-Gel-Peeling-So-Good-Perawatan-Wajah-Eksfoliasi-Ringan-untuk-Sensitive-Skin-Scrub-Wajah-AHA-BHA-PHA-dan-Centella-Asiatica-i.255365082.20571831223",
     32500, 4.9, 63000),
    ("SOMETHINC Serum Series (Glass-C / Niacinamide / Retinol / Hyaluronic9)",
     "https://shopee.co.id/SOMETHINC-Serum-Series-Glass-C-12-Vitamin-C-Gentle-Bright-Niacinamide-Advanced-Peeling-20-AHA-BHA-Level-1-Retinol-Calm-Down-Skinpair-Barrier-Diamond-Glow-Phyto-Hyaluronic9-Advanced-Wajah-Cerah-No-Dark-Spot-Age-Spot-i.195455930.43663200498",
     119000, 5.0, 11200),
    ("AVOSKIN Miraculous Refining Serum 30ml",
     "https://shopee.co.id/AVOSKIN-Miraculous-Refining-Serum-30ml-i.68111.2833773785",
     259300, 4.9, 1000),
    ("SOMETHINC Skinglow Bright Radiance Moisturizer",
     "https://shopee.co.id/SOMETHINC-SKINGLOW%C2%AE-BRIGHT-RADIANCE-MOISTURIZER-Pelembab-Mencerahkan-Memudarkan-Noda-Hitam-i.195455930.12923585738",
     100000, 4.9, 4900),
    ("Wardah Sunscreen UV Shield All Series",
     "https://shopee.co.id/Wardah-Sunscreen-UV-Shield-All-Series-(Airy-Smooth-Essential-Gel-Acne-Calming-Aqua-Fresh-Tone-Up-Physical)-Sunscreen-SPF-50-PA-SPF-35-PA-Sunscreen-No.1-Sunscreen-of-the-Year-Skincare-i.59763733.28218154859",
     40180, 4.9, 249900),
    ("OMG Oh My Glam Coverlast Liquid Foundation",
     "https://shopee.co.id/-NEW-SHADE!-OMG-Oh-My-Glam-Coverlast-Liquid-Foundation-Flawless-Coverage-Tahan-12-Jam-i.401724234.14776749929",
     25900, 4.9, 106700),
    ("OMG Oh My Glam Coverlast Two Way Cake Powder",
     "https://shopee.co.id/-BELIA-OMG-Oh-My-Glam-Coverlast-Two-Way-Cake-Powder-Bedak-Padat-Kontrol-Minyak-Hingga-8-Jam-TWC-i.48382819.22613655801",
     27100, 4.9, 102100),
    ("WARDAH TWC Colorfit Velvet Powder Foundation 11g",
     "https://shopee.co.id/NEW!-WARDAH-TWC-Colorfit-Velvet-Powder-Foundation-11-g-Two-Way-Cake-No-1-di-Indonesia-Medium-to-High-Coverage-Bedak-Padat-Tahan-Lama-16-Jam-Flawless-Longlasting-for-All-Skin-Type-Breathable-Soft-Blur-Cover-TWC-Cushion-Foundation-Matte-Makeup-i.59763733.5680847884",
     61570, 4.9, 109300),
    ("OMG Oh My Glam Lip Series Mattelast Lip Cream 2.9g",
     "https://shopee.co.id/-BELIA-OMG-OH-MY-GLAM-Lip-Series-Mattelast-Lip-Cream-2.9-g-Lip-Tint-Lipstick-i.48382819.5183050900",
     23700, 4.9, 536600),
    ("HANASUI Mattedorable Lipcream Boba Edition",
     "https://shopee.co.id/HANASUI-(BOBBA)-Mattedorable-Lipcream-Boba-Edition-i.5139506.6270029186",
     14700, 4.8, 65700),
    ("MAYBELLINE Superstay Matte Ink",
     "https://shopee.co.id/-SHOPEE-BEAUTY-AWARDS-BEST-LIP-MAKEUP-MAYBELLINE-OFFICIAL-Superstay-Matte-Ink-Liquid-Long-Lasting-Waterproof-Matte-Lipstick-Lipcream-Make-Up-Transferproof-Tahan-16-Jam-Vinyl-Ink-Lipstik-Lip-Gloss-Liptin-i.62582411.1835009883",
     85000, 4.9, 3100000),
    ("IMPLORA Urban Lip Cream Matte",
     "https://shopee.co.id/-BEST-SELLER-IMPLORA-OFFICIAL-Urban-Lip-Cream-Matte-20-Shades-Long-lasting-Formula-Ringan-dan-Lembut-Waterproof-Tahan-Lama-Matte-Finish-Ringan-Auto-Set-2-Menit-Skincare-Infused-Vitamin-E-Transferproof-Pigmented-Liquid-Lipstick-Anti-Luntur-i.10960132.1993631370",
     23000, 4.9, 132900),
    ("SH-RD Protein Cream 50ml Vitamin Rambut",
     "https://shopee.co.id/SH-RD-Protein-Cream-(50ml-)-Moisturizer-Vitamin-Rambut-Hair-Care-Rambut-Ngembang-Kering-Kusut-i.167106407.3606884214",
     219000, 4.9, 85900),
    ("Lab On Hair Anti Hair Fall Ampoule Serum 60ml",
     "https://shopee.co.id/Lab-On-Hair-Anti-Hair-Fall-Ampoule-Serum-Hair-Oil-Serum-60-ml-Vitamin-Rambut-Growth-Rambut-Anti-Rontok-Kering-Bercabang-Mempercepat-Pertumbuhan-Rambut-Penumbuh-Rambut-Water-Base-Minyak-Rambut-dengan-PDO-Biotin-Ginger-Hair-Serum-Hair-Ampoule-Serum-i.811333308.19620542816",
     177900, 4.9, 74400),
    ("Makarizo Hair Energy Scentsations Hair Fragrance",
     "https://shopee.co.id/Makarizo-Hair-Energy-Scentsations-Hair-Fragrance-30ml-100ml-Parfum-Rambut-Vitamin-UV-Protection-i.327796207.12236741538",
     12899, 4.8, 13200),
    ("Pantene Pro-V Perfect On Conditioner 40ml",
     "https://shopee.co.id/Pantene-Pro-V-Perfect-On-Conditioner-Tanpa-Bilas-40ml-i.30736001.3177767098",
     20700, 4.9, 12600),
    ("Grace and Glow Brightening Body Wash 400ml",
     "https://shopee.co.id/Grace-and-Glow-Brightening-Body-Wash-400ml-Sabun-Mandi-Cair-Mencerahkan-Kulit-dan-Melembabkan-Kulit-with-Niacinamide-dan-Glutathione-dengan-Wangi-Parfum-Mewah-Body-Soap-Shower-Gel-Wangi-Kemasan-Botol-dan-Refill-Pouch-Isi-Ulang-i.542670804.2994521472",
     58900, 4.9, 152200),
    ("Rintik Skincare No Bump Lotion 100ml",
     "https://shopee.co.id/Rintik-Skincare-No-Bump-Lotion-100ml-Bantu-Hilangkan-Kulit-Bruntusan-Chicken-Skin-Jerawat-Punggung-Pantat-Budukan-BPOM-Halal-i.513172314.19069132449",
     92102, 4.9, 72300),
    ("Vaseline Healthy Bright UV Extra Brightening Body Lotion 400ml",
     "https://shopee.co.id/Vaseline-Healthy-Bright-UV-Extra-Brightening-No-1-Body-Lotion-untuk-Kulit-Cerah-Merata-Plumpy-Terlindungi-dengan-GlutaGlow-Pro-Ceramide-Triple-UV-Filters-400mL-i.1882916412.46964190536",
     68200, 4.8, 80900),
    ("CITRA Body Serum Radiant Fresh Bio-AHA 100ml",
     "https://shopee.co.id/CITRA-BODY-SERUM-RADIANT-FRESH-BIO-AHA-100ML-i.1891521336.58065929194",
     17200, 4.8, 207300),
    ("SCARLETT Fragrance Brightening Body Lotion 300ml",
     "https://shopee.co.id/SCARLETT-Fragrance-Brightening-Body-Lotion-Bright-Mood-Booster-300ml-4-Essential-Oils-3X-Brightening-Power-Solusi-Untuk-Kulit-Kering-Menenangkan-Mencerahkan-Kulit-Wangi-Tahan-Seharian-i.255365082.25631959698",
     59000, 4.9, 138100),
    ("Marina UV White Healthy & Glow Handbody Lotion 460ml",
     "https://shopee.co.id/Handbody-Lotion-Marina-UV-White-Healthy-Glow-460-mL-i.18163317.1195941659",
     22700, 4.9, 6400),
]


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Produk"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="C0392B")
    need_fill = PatternFill("solid", fgColor="FDEBD0")  # kolom wajib diisi Kan

    for col, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for r, (nama, url, harga, rating, ulasan) in enumerate(PRODUCTS, start=2):
        ws.cell(row=r, column=1, value=nama)
        ws.cell(row=r, column=2, value="shopee")
        url_cell = ws.cell(row=r, column=3, value=url)
        url_cell.hyperlink = url
        url_cell.font = Font(color="0563C1", underline="single")
        ws.cell(row=r, column=4, value=harga).number_format = '#,##0'
        ws.cell(row=r, column=5, value=rating)
        ws.cell(row=r, column=6, value=ulasan).number_format = '#,##0'
        for col in (7, 8, 9):  # link_affiliate, url_gambar, deskripsi
            ws.cell(row=r, column=col).fill = need_fill

    widths = [55, 12, 18, 14, 10, 16, 18, 18, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_properties.pageSetUpPr = None

    out = os.path.expanduser("~/workspace/your_files/katalog-cekdulu-beauty.xlsx")
    wb.save(out)
    print(f"OK: {out} ({len(PRODUCTS)} produk)")


if __name__ == "__main__":
    main()
